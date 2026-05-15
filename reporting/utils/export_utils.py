import csv
import io
from datetime import datetime
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT
from xml.sax.saxutils import escape

# COULEURS SOMATRIN
COLOR_PRIMARY = "E87722"
COLOR_DARK = "1a2c4e"
COLOR_LIGHT_GRAY = "f9f9f9"
COLOR_BORDER = "e0e3e7"


def _hex(c):
    """ReportLab HexColor attend un préfixe #."""
    return "#" + c if not str(c).startswith("#") else c


# ═══════════════════════════════════════════════════════════════════════════
# EXPORT EXCEL
# ═══════════════════════════════════════════════════════════════════════════

def export_to_excel(data_list, filename, column_names, title="SOMATRIN Dashboard"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Données"

    header_fill = PatternFill(start_color=COLOR_DARK, end_color=COLOR_DARK, fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    title_font = Font(bold=True, size=14, color=COLOR_DARK)
    border = Border(
        left=Side(style='thin', color=COLOR_BORDER),
        right=Side(style='thin', color=COLOR_BORDER),
        top=Side(style='thin', color=COLOR_BORDER),
        bottom=Side(style='thin', color=COLOR_BORDER)
    )

    # Titre
    ws.merge_cells('A1:' + chr(64 + len(column_names)) + '1')
    title_cell = ws['A1']
    title_cell.value = title
    title_cell.font = title_font
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 25

    # Date
    ws.merge_cells('A2:' + chr(64 + len(column_names)) + '2')
    date_cell = ws['A2']
    date_cell.value = f"Généré le: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    date_cell.alignment = Alignment(horizontal='center')
    ws.row_dimensions[2].height = 15

    # Headers
    for col_idx, col_name in enumerate(column_names, start=1):
        cell = ws.cell(row=4, column=col_idx)
        cell.value = col_name
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border

    ws.row_dimensions[4].height = 20

    # Données
    for row_idx, data_row in enumerate(data_list, start=5):
        for col_idx, col_name in enumerate(column_names, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)

            if isinstance(data_row, dict):
                value = data_row.get(col_name, '')
            else:
                value = data_row[col_idx - 1] if col_idx - 1 < len(data_row) else ''

            cell.value = value
            cell.border = border

            if isinstance(value, (int, float)) and 'montant' in col_name.lower():
                cell.font = Font(color=COLOR_PRIMARY, bold=True)
                cell.number_format = '#,##0.00'

            cell.alignment = Alignment(horizontal='center' if isinstance(value, (int, float)) else 'left')

    # Largeur colonnes
    for col_idx, col_name in enumerate(column_names, start=1):
        ws.column_dimensions[chr(64 + col_idx)].width = max(len(col_name) + 2, 15)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


# ═══════════════════════════════════════════════════════════════════════════
# EXPORT CSV
# ═══════════════════════════════════════════════════════════════════════════

def export_to_csv(data_list, filename, column_names):
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response.write('\ufeff')

    writer = csv.DictWriter(response, fieldnames=column_names)
    writer.writeheader()

    for data_row in data_list:
        if isinstance(data_row, dict):
            row = {col: data_row.get(col, '') for col in column_names}
            writer.writerow(row)
        else:
            row = {column_names[i]: data_row[i] for i in range(len(column_names))}
            writer.writerow(row)

    return response


# ═══════════════════════════════════════════════════════════════════════════
# EXPORT PDF
# ═══════════════════════════════════════════════════════════════════════════

def export_to_pdf(data_list, filename, column_names, title="SOMATRIN Dashboard"):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=15,
        leftMargin=15,
        topMargin=30,
        bottomMargin=30,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=16,
        textColor=colors.HexColor(_hex(COLOR_DARK)),
        spaceAfter=6,
        alignment=1,
    )

    elements = []

    elements.append(Paragraph(title, title_style))
    elements.append(Paragraph(
        f"<font size=9 color='gray'>Généré le {datetime.now().strftime('%d/%m/%Y %H:%M')}</font>",
        styles['Normal']
    ))
    elements.append(Spacer(1, 0.3 * inch))

    table_data = [column_names]

    for data_row in data_list:
        row = []
        for col_name in column_names:
            if isinstance(data_row, dict):
                value = data_row.get(col_name, '')
            else:
                col_idx = column_names.index(col_name)
                value = data_row[col_idx] if col_idx < len(data_row) else ''

            if isinstance(value, float):
                row.append(f"{value:,.2f}")
            else:
                row.append(str(value))

        table_data.append(row)

    table = Table(table_data)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(_hex(COLOR_DARK))),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 11),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('ALIGN', (0, 1), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor(_hex(COLOR_LIGHT_GRAY))]),
        ('GRID', (0, 0), (-1, -1), 1, colors.HexColor(_hex(COLOR_BORDER))),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
    ]))

    elements.append(table)
    elements.append(Spacer(1, 0.3 * inch))
    elements.append(Paragraph(
        "<font size=8 color='gray'>SOMATRIN Dashboard | Export © 2026</font>",
        styles['Normal']
    ))

    doc.build(elements)

    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    return response


# Colonnes export Pointages (libellé Excel/PDF = clé dict)
POINTAGES_EXPORT_COLUMNS = [
    "Date",
    "Opération",
    "Site",
    "Société",
    "Engin / Affectation",
    "Équipe",
    "Ouvrage",
    "Foration",
    "ML estimés",
    "Statut",
]


def pointages_rows_to_export_dicts(rows):
    """Convertit les lignes métier (clés anglaises) en dicts pour export Excel/CSV/PDF."""
    out = []
    for r in rows or []:
        try:
            ml = float(r.get("ml_estime") or 0)
        except (TypeError, ValueError):
            ml = 0.0
        out.append({
            "Date": r.get("date") or "—",
            "Opération": r.get("operation") or "—",
            "Site": r.get("site") or "—",
            "Société": r.get("societe") or "—",
            "Engin / Affectation": r.get("engin") or "—",
            "Équipe": r.get("equipe") or "—",
            "Ouvrage": r.get("ouvrage") or "—",
            "Foration": r.get("foration") or "—",
            "ML estimés": round(ml, 1) if ml else 0.0,
            "Statut": r.get("statut") or "—",
        })
    return out


def export_pointages_operations_pdf(
    data_list,
    filename,
    title="SOMATRIN — Pointages opérations & foration",
    subtitle_lines=None,
    source_note=None,
):
    """
    PDF paysage : bandeau titre, filtres, tableau colonnes dimensionnées, retours à la ligne (Paragraph).
    """
    subtitle_lines = subtitle_lines or []
    source_note = source_note or ""

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=18,
        leftMargin=18,
        topMargin=42,
        bottomMargin=36,
    )
    page_w = landscape(A4)[0] - 36

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "PtTitle",
        parent=styles["Heading1"],
        fontSize=17,
        leading=22,
        textColor=colors.HexColor(_hex(COLOR_DARK)),
        spaceAfter=4,
        alignment=TA_LEFT,
    )
    sub_style = ParagraphStyle(
        "PtSub",
        parent=styles["Normal"],
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#4b5563"),
        alignment=TA_LEFT,
        spaceAfter=2,
    )
    note_style = ParagraphStyle(
        "PtNote",
        parent=styles["Normal"],
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#6b7280"),
        alignment=TA_LEFT,
        spaceAfter=10,
    )
    hdr_style = ParagraphStyle(
        "PtHdr",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=7.5,
        leading=10,
        textColor=colors.white,
        alignment=TA_CENTER,
    )
    cell_left = ParagraphStyle(
        "PtCellL",
        parent=styles["Normal"],
        fontSize=7,
        leading=9,
        alignment=TA_LEFT,
        wordWrap="CJK",
    )
    cell_center = ParagraphStyle(
        "PtCellC",
        parent=styles["Normal"],
        fontSize=7,
        leading=9,
        alignment=TA_CENTER,
    )
    cell_right = ParagraphStyle(
        "PtCellR",
        parent=styles["Normal"],
        fontSize=7,
        leading=9,
        alignment=TA_RIGHT,
    )

    def _p(text, style):
        s = escape(str(text if text is not None else ""))
        return Paragraph(s.replace("\n", "<br/>"), style)

    elements = []
    elements.append(_p(title, title_style))
    elements.append(
        _p(
            f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}",
            sub_style,
        )
    )
    for line in subtitle_lines:
        if line:
            elements.append(_p(line, sub_style))
    if source_note:
        elements.append(_p(source_note, note_style))
    elements.append(Spacer(1, 0.12 * inch))

    col_keys = POINTAGES_EXPORT_COLUMNS
    # Largeurs (pt) — total ≈ page_w
    col_widths = [52, 76, 70, 58, 86, 72, 248, 40, 48, 62]
    scale = page_w / sum(col_widths)
    col_widths = [w * scale for w in col_widths]

    header_row = [_p(h, hdr_style) for h in col_keys]
    table_data = [header_row]

    for data_row in data_list:
        row = []
        for col_name in col_keys:
            if isinstance(data_row, dict):
                raw = data_row.get(col_name, "")
            else:
                raw = ""
            if col_name == "ML estimés":
                row.append(_p(raw, cell_right))
            elif col_name in ("Foration", "Statut", "Date"):
                row.append(_p(raw, cell_center))
            else:
                row.append(_p(raw, cell_left))
        table_data.append(row)

    table = Table(table_data, colWidths=col_widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(_hex(COLOR_DARK))),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ALIGN", (0, 0), (-1, 0), "CENTER"),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor(_hex(COLOR_LIGHT_GRAY))],
                ),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor(_hex(COLOR_BORDER))),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("LINEBELOW", (0, 0), (-1, 0), 2, colors.HexColor(_hex(COLOR_PRIMARY))),
            ]
        )
    )
    elements.append(table)
    elements.append(Spacer(1, 0.22 * inch))
    elements.append(
        Paragraph(
            f"<font size=8 color='#9ca3af'>SOMATRIN — Export confidentiel · "
            f"{len(data_list)} ligne(s) · © {datetime.now().year}</font>",
            styles["Normal"],
        )
    )

    doc.build(elements)
    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
